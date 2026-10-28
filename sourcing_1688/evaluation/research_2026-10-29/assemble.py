"""Сводная таблица всех позиций действующих заказов (28.10): данные для шкалы v2."""
import glob, json, os, re, sys
import openpyxl
P = "/Volumes/ADATA SD810/WB/Подбор"; W = f"{P}/wb_slice_2026-10-22"; SP = "/private/tmp/claude-501/-Volumes-ADATA-SD810-WB/a25f6ca6-5448-40ba-9e30-e0ec147f4619/scratchpad"
pool = {p["id"]: p for p in json.load(open(f"{W}/pool_1688.json"))}
kq = {x["id"]: x for x in json.load(open(f"{W}/kq_sets.json"))}
det = json.load(open(f"{W}/details.json")); pr = json.load(open(f"{W}/photo_review.json"))
# 28.10: карточки, восстановленные из черновика после обрыва details.json, заменяются перекачанными (refetch.py → details_refetch.json)
_rf = os.path.join(os.path.dirname(os.path.abspath(__file__)), "details_refetch.json")
for _rf in (_rf, _rf.replace("details_refetch.json", "details_refetch2.json")):   # два потока перекачки
  if os.path.exists(_rf):
    for _k, _v in json.load(open(_rf)).items():
        if det.get(_k, {}).get("restored") or _k not in det: det[_k] = _v
taste = {}
for f in ["kimi_taste.json", "kimi_taste3.json", "kimi_taste4.json", "kimi_taste5.json", "codex_taste6.json", "codex_taste6a.json", "codex_taste6b.json"] + sorted(glob.glob(f"{W}/codex_taste_w*.json")):
    f = f if f.startswith("/") else f"{W}/{f}"
    if os.path.exists(f):
        for k, v in json.load(open(f)).items():
            taste.setdefault(k, v)
srev = json.load(open(f"{W}/sets_review.json")); kqr = json.load(open(f"{W}/kq_sets_review.json"))
wsrc = {}
for d in [f"{P}/Черновик_заказа_бижутерия_2026-10-22.xlsx", f"{SP}/draft_sets3.xlsx"]:
    for r in list(openpyxl.load_workbook(d, read_only=True)["Черновик заказа"].iter_rows(values_only=True))[1:]:
        wsrc.setdefault(re.search(r"offer/(\d+)", r[4]).group(1), r[7])
def img(k):
    for p in [pool.get(k, {}).get("img") if str(pool.get(k, {}).get("img", "")).startswith("/") else None, f"{W}/sets7/{k}.jpg", f"{W}/kqsets/{k}.jpg", f"{W}/wide_img/{k}.jpg", f"{W}/codex6/{k}.jpg", f"{W}/t3/{k}.jpg"] + [f"{W}/full/{k}_{i}.jpg" for i in range(3)] + glob.glob(f"{W}/img/*/{k}*") + glob.glob(f"{P}/1688_данные/novinki55_2026-10-21/cand/*/{k}.jpg"):
        if p and os.path.exists(p) and not os.path.basename(p).startswith("._"):
            return p
oid = lambda u: re.search(r"offer/(\d+)", str(u)).group(1)
items = []
for order, f in [("100", "Заказ_бижутерия_100_2026-10-23.xlsx"), ("200", "Заказ_бижутерия_200_баланс_2026-10-24.xlsx"), ("1000-ч1", "Заказ_1000_часть1_одиночные_885_без_дублей_2026-10-27.xlsx"), ("добор", "Заказ_комплекты_добор_185_2026-10-27.xlsx")]:
    for r in openpyxl.load_workbook(f"{P}/{f}", read_only=True).worksheets[0].iter_rows(min_row=2, values_only=True):
        if not isinstance(r[0], int):
            continue
        k = oid(r[5]); t = taste.get(k, {})
        items.append(dict(order=order, row=r[0], group=r[2], src=r[3], seg=r[4], ids=[k], cny=r[6], w=r[7], wsrc=wsrc.get(k) or ("карточка" if "重量" in (det.get(k, {}).get("weights") or "") else "допущение"),
                          P=r[9], psrc=r[10], roi=r[11], roi1330=r[12], qty=r[13], note=r[15] or "", title=(pool.get(k) or kq.get(k) or {}).get("title", ""),
                          pv=pr.get(k, {}).get("v", ""), taste=t.get("taste"), why=t.get("why", ""), like=t.get("like_wb"), det=det.get(k, {}), imgs=[img(k)], url=r[5]))
for r in openpyxl.load_workbook(f"{P}/Заказ_1000_часть2_комплекты_115_без_дублей_2026-10-27.xlsx", read_only=True).worksheets[0].iter_rows(min_row=2, values_only=True):
    if not isinstance(r[0], int):
        continue
    a = oid(r[5]); b = oid(r[7]) if r[7] and "offer" in str(r[7]) else None
    seg = "А2/Б колье|подвеска+серьги" if b else (kq.get(a, {}).get("seg") or "набор прочий")
    rv = srev.get(f"{a}+{b}", {}).get("v", "") if b else kqr.get(a, {}).get("v", "")
    items.append(dict(order="1000-ч2", row=r[0], group=str(r[3]), src=("сборный" if b else "готовый"), seg=seg, ids=[a] + ([b] if b else []), cny=r[9], w=r[10], wsrc=r[11], P=1330, psrc="верхний ярус 1 330 ₽",
                      roi=r[12], roi1330=r[12], qty=r[15], note=r[17] or "", title=" | ".join((pool.get(x) or kq.get(x) or {}).get("title", "") for x in [a] + ([b] if b else [])),
                      pv=rv, taste=None, why=r[17] or "", like=None, det=det.get(a, {}), imgs=[img(a)] + ([img(b)] if b else []), url=r[5], url2=r[7] if b else None, pair_score=r[14]))
json.dump(items, open("items.json", "w"), ensure_ascii=False)
from collections import Counter
print(len(items), Counter(i["order"] for i in items), "без фото", sum(1 for i in items if not all(i["imgs"])), "без названия", sum(1 for i in items if not i["title"]), "без pv", sum(1 for i in items if not i["pv"]))
