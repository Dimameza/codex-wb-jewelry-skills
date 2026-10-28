"""Заказ комплектов «колье + серьги» (24.10): пары собраны из товаров разных фабрик 1688 (решение владельца — комплектуем сами).
Пары предложил Codex по фото, каждую пару проверил Claude по фото (sets_review.json). Цена продажи комплекта — верхний ярус владельца 1 330 ₽.
ROI = (0,0343·P − 6,30) / ((22 + 0,285·вес_колье+серёг)/12,72 + ¥колье + ¥серёг) − 1."""
import json, os, re, statistics as st, sys
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XImage
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image

W = os.path.dirname(os.path.abspath(__file__))
N = int(sys.argv[1]); OUT = sys.argv[2]; P = 1330; QTY = int(os.environ.get("QTY", "50"))   # штук на комплект
pool = {p["id"]: p for p in json.load(open(f"{W}/pool_1688.json"))}
det = json.load(open(f"{W}/details.json"))
pairs = json.load(open(f"{W}/sets7_pairs.json"))
rev = json.load(open(f"{W}/sets_review.json")) if os.path.exists(f"{W}/sets_review.json") else {}
WDEF = {"колье": 30, "чокер": 30, "подвеска": 15, "бусы": 40, "серьги": 15}


def nums(s):
    return [float(x) for x in re.findall(r"\d+(?:\.\d+)?", s or "")]


def weight(k):
    t = det.get(k, {}).get("weights") or ""
    ws = []
    if "重量" in t:
        for r in t.split("\n")[1:]:
            if r.startswith("商品详情"):
                break
            n = nums(r)
            if n:
                ws.append(n[-1])
    ws = [x for x in ws if 2 < x <= 300]
    return (st.median(ws), "карточка") if ws else (WDEF.get(pool[k]["seg"], 30), "допущение")


def cny(k):
    n = nums(det.get(k, {}).get("sku_price_scale"))
    return max(n) if n else pool[k]["price"]


def roi(p, y, w):
    return ((0.0343 * p - 6.30) / ((22 + 0.285 * w) / 12.72 + y) - 1) * 100


# офферы прошлых заказов: готовый комплект из них не повторяем (env PREV_ORDER через «|»)
import openpyxl
PREV = set()
for prev in filter(None, os.environ.get("PREV_ORDER", "").split("|")):
    for x in list(openpyxl.load_workbook(prev, read_only=True).worksheets[0].iter_rows(values_only=True))[1:]:
        if isinstance(x[0], int):
            PREV.add(re.search(r"offer/(\d+)", x[5]).group(1))
SEX = json.load(open(f"{W}/sets_exclude.json")) if os.path.exists(f"{W}/sets_exclude.json") else {}   # дубли моделей, найденные глазами
rows, used = [], set()
for p in pairs:
    a, b = str(p["neck"]), str(p["ear"])
    key = f"{a}+{b}"
    v = rev.get(key, {}).get("v", "")
    if not v.startswith("ок") or a in used or b in used or a not in pool or b not in pool or a in SEX or b in SEX:   # 27.10: SEX и для сборных — дубль модели колье
        continue
    ya, yb = cny(a), cny(b); (wa, sa), (wb_, sb) = weight(a), weight(b)
    r = roi(P, ya + yb, wa + wb_)
    if r < 100:
        continue
    used |= {a, b}
    rows.append({"a": a, "b": b, "ya": ya, "yb": yb, "w": wa + wb_, "wsrc": "карточка" if sa == sb == "карточка" else "допущение",
                 "roi": round(r), "roi1600": round(roi(1600, ya + yb, wa + wb_)), "score": p.get("score"), "color": p.get("color"), "motif": p.get("motif"),
                 "why": (p.get("why") or "") + ("".join(f" · {nm} уже в прошлом заказе отдельно" for nm, x in (("колье", a), ("серьги", b)) if x in PREV)), "rv": v, "rc": rev[key].get("c", ""),
                 "tier": 0 if r >= 200 and v != "ок, условно" else 1 if v != "ок, условно" else 2})
# готовые комплекты одной фабрики (поиск 1688 24.10, проверены Claude по фото: kq_sets_review.json)
kq = {x["id"]: x for x in json.load(open(f"{W}/kq_sets.json"))} if os.path.exists(f"{W}/kq_sets.json") else {}
kqr = json.load(open(f"{W}/kq_sets_review.json")) if os.path.exists(f"{W}/kq_sets_review.json") else {}
SCORE = {"ок, сильно": 5, "ок": 4, "ок, условно": 3}
for k, rv_ in kqr.items():
    if os.environ.get("NECK_ONLY") and not any(w in kq[k]["title"] for w in ("项链", "项圈", "颈链", "锁骨链", "吊坠", "毛衣链")):
        continue   # 26.10: основа — колье и комплекты; набор без колье (браслет/кольца) не берём — браслеты зимой не берут
    if k in used or k in PREV or k in SEX or k not in det or det[k].get("blocked"):
        continue
    n = nums(det[k].get("sku_price_scale")); y = max(n) if n else kq[k]["price"]
    t = det[k].get("weights") or ""; wl = []
    if "重量" in t:
        for ln in t.split("\n")[1:]:
            if ln.startswith("商品详情"):
                break
            nn = nums(ln)
            if nn:
                wl.append(nn[-1])
    wl = [x for x in wl if 2 < x <= 300]
    if not wl:   # вес строкой в атрибутах карточки: «重量\n20» или «重量\n20g» (26.10)
        m = re.search(r"重量\s*\n\s*(\d+(?:\.\d+)?)\s*g?\s*\n", det[k].get("colors") or "")
        if m and 2 < float(m.group(1)) <= 300:
            wl = [float(m.group(1))]
    w, ws_ = (st.median(wl), "карточка") if wl else (60, "допущение")
    r = roi(P, y, w)
    if r < 100:
        continue
    used.add(k)
    rows.append({"a": k, "b": None, "ya": y, "yb": 0, "w": w, "wsrc": ws_, "roi": round(r), "roi1600": round(roi(1600, y, w)), "score": SCORE[rv_["v"]],
                 "color": "", "motif": kq[k]["seg"], "why": "готовый комплект одной фабрики: " + kq[k]["title"][:50], "rv": rv_["v"], "rc": "",
                 "tier": 0 if r >= 200 and rv_["v"] != "ок, условно" else 1 if rv_["v"] != "ок, условно" else 2})
rows.sort(key=lambda r: (r["tier"], -(r["score"] or 0), -r["roi"]))
rows = rows[:N]
TIER = {0: "брать (ROI ≥200 % при 1 330 ₽)", 1: "брать (ROI 100–200 %)", 2: "условно (фото на грани)"}
FILL = {0: "C6EFCE", 1: "D9EAD3", 2: "FFEB9C"}
IMG = lambda k: k and next((f for f in (f"{W}/sets7/{k}.jpg", f"{W}/full/{k}_0.jpg", f"{W}/kqsets/{k}.jpg") if os.path.exists(f)), None)
wbk = Workbook(); ws = wbk.active; ws.title = f"Комплекты {len(rows)}"
H = ["№", "фото колье / комплекта", "фото серёг", "группа", "цвет / мотив", "колье или комплект 1688", "¥ колье / комплекта", "серьги 1688 (для сборного)", "¥ серёг", "¥ комплект", "вес, г", "вес — откуда",
     "ROI % при 1 330 ₽", "ROI % при 1 600 ₽", "оценка пары", "кол-во", "сумма ¥", "почему комплект"]
ws.append(H)
for c in ws[1]:
    c.font = Font(bold=True); c.alignment = Alignment(wrap_text=True, vertical="top")
for i, wd in enumerate([4, 15, 15, 16, 14, 15, 7, 15, 7, 8, 7, 10, 8, 8, 7, 7, 9, 70]):
    ws.column_dimensions[chr(65 + i)].width = wd
tot = 0
for n, r in enumerate(rows, 1):
    y = r["ya"] + r["yb"]; tot += QTY * y
    ws.append([n, "", "", TIER[r["tier"]], f"{r['color'] or ''} / {r['motif'] or ''}", f"https://detail.1688.com/offer/{r['a']}.html", r["ya"],
               (f"https://detail.1688.com/offer/{r['b']}.html" if r["b"] else "— готовый комплект"), r["yb"], round(y, 2), round(r["w"]), r["wsrc"], r["roi"], r["roi1600"], r["score"], QTY,
               round(QTY * y, 1), (r["why"] or "") + (f" · проверка: {r['rc']}" if r["rc"] else "")])
    row = ws.max_row; ws.row_dimensions[row].height = 84
    for c in ws[row]:
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(row, 4).value = ("сборный · " if r["b"] else "готовый · ") + TIER[r["tier"]]
    ws.cell(row, 4).fill = PatternFill("solid", fgColor=FILL[r["tier"]])
    for col, k in (("B", r["a"]), ("C", r["b"])):
        f = IMG(k)
        if f:
            tmp = f"/tmp/set_{k}.png"; im = Image.open(f).convert("RGB"); im.thumbnail((105, 105)); im.save(tmp); ws.add_image(XImage(tmp), f"{col}{row}")
ws.append([]); ws.append(["", "", "", "ИТОГО", f"{len(rows)} комплектов × {QTY}", "", "", "", "", "", "", "", "", "", "", QTY * len(rows), round(tot),
                          f"≈ {round(tot * 12.72 / 1000)} тыс ₽ закупки без карго и без сборки комплектов на ФФ"])
ws.freeze_panes = "D2"
m = wbk.create_sheet("Метод")
for line in [__doc__,
             "Правила пары: один металл и оттенок, одна стилистика и мотив, совместимый масштаб; каждый оффер — не больше чем в одном комплекте.",
             "Не учтено в ROI: упаковка комплекта (одна карточка/пакет на два товара) и работа ФФ по комплектации.",
             "Перед оплатой: выбрать цвет так, чтобы колье и серьги совпали по оттенку; уточнить вес, где «допущение»."]:
    m.append([line])
m.column_dimensions["A"].width = 150
wbk.save(OUT)
from collections import Counter
print(OUT, len(rows), dict(Counter(TIER[r["tier"]] for r in rows)), "сборных", sum(1 for r in rows if r["b"]), "готовых", sum(1 for r in rows if not r["b"]), "¥", round(tot))
