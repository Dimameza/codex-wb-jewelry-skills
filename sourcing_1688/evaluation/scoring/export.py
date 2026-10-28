import json, os
from collections import Counter
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XImage
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image, ImageFile
ImageFile.LOAD_TRUNCATED_IMAGES = True
rows = json.load(open("scored.json")); OUT = "/Volumes/ADATA SD810/WB/Подбор/Оценка_по_критериям_v2_все_1485_проверено_глазами_2026-10-28.xlsx"
K = ["спрос", "читаемость", "профиль", "повод", "качество", "комплектность"]; MAX = dict(спрос=20, читаемость=20, профиль=20, повод=15, качество=15, комплектность=10)
FILL = {"оплатить": "C6EFCE", "тест 20–30 шт": "FFEB9C"}
wb = Workbook(); ws = wb.active; ws.title = "Баллы"
H = ["фото", "заказ", "№ в заказе", "сегмент", "основной стиль (3 описания)", "все стили (≥2 из 3)", "описание Claude", "описание Codex", "описание Kimi", "оффер 1688", "¥", "вес, г", "ИТОГ (из 100)", "решение по v2 (честная цена)", "решение, если 1 330 ₽ подтвердится",
     "экономика v2", "честный ROI %", "нужная цена для ROI 100 %, ₽", "экономика — расчёт"]
for k in K: H += [f"{k} ({MAX[k]})", f"{k} — почему"]
H += ["стоп-факторы", "описание фото (Codex/Kimi)", "название 1688"]
ws.append(H)
for c in ws[1]: c.font = Font(bold=True); c.alignment = Alignment(wrap_text=True, vertical="top")
widths = [14, 8, 6, 14, 18, 22, 40, 40, 40, 14, 6, 6, 8, 16, 16, 11, 8, 10, 40] + [7, 45] * 6 + [14, 40, 30]
from openpyxl.utils import get_column_letter
for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
for r in rows:
    line = ["", r["order"], r["row"], r["seg"], r["style"], r.get("style_all", ""), r["descr"].get("claude", ""), r["descr"].get("codex", ""), r["descr"].get("kimi", ""), r["url"], r["cny"], r["w"], r["total"], r["decision"], r["decision1330"], r["econ"], r["eroi"], r.get("need_P"), r["etxt"]]
    for k in K: line += [r["R"][k][0], r["R"][k][1]]
    line += ["; ".join(r["stops"]), r.get("why", ""), r.get("title", "")]
    ws.append(line); n = ws.max_row; ws.row_dimensions[n].height = 72
    for c in ws[n]: c.alignment = Alignment(wrap_text=True, vertical="top")
    for col, key in ((14, r["decision"]), (15, r["decision1330"])):
        if key in FILL: ws.cell(n, col).fill = PatternFill("solid", fgColor=FILL[key])
    p = r["imgs"][0]
    if p and os.path.exists(p):
        tmp = f"/tmp/sc_{r['order']}_{r['row']}.png"; im = Image.open(p).convert("RGB"); im.thumbnail((90, 90)); im.save(tmp); ws.add_image(XImage(tmp), f"A{n}")
ws.freeze_panes = "F2"; ws.auto_filter.ref = ws.dimensions
s = wb.create_sheet("Сводка")
s.append(["заказ", "позиций", "средний балл", "≥70 баллов", "55–69", "<55", "экономика v2 проходит", "если 1 330: оплатить", "если 1 330: тест", "если 1 330: не брать"])
for o in ["100", "200", "1000-ч1", "1000-ч2", "добор", "все"]:
    R = [r for r in rows if o == "все" or r["order"] == o]; d = Counter(r["decision1330"] for r in R)
    s.append([o, len(R), round(sum(r["total"] for r in R) / len(R), 1), sum(r["total"] >= 70 for r in R), sum(55 <= r["total"] < 70 for r in R), sum(r["total"] < 55 for r in R),
              sum(r["econ"] == "проходит" for r in R), d["оплатить"], d["тест 20–30 шт"], sum(v for k, v in d.items() if k.startswith("не брать"))])
s.append([]); s.append(["критерий", "распределение баллов (балл: позиций)"])
for k in K: s.append([f"{k} (макс {MAX[k]})", ", ".join(f"{b}: {c}" for b, c in sorted(Counter(r["R"][k][0] for r in rows).items(), reverse=True))])
for c in s[1]: c.font = Font(bold=True)
s.column_dimensions["A"].width = 26; s.column_dimensions["B"].width = 60
t = wb.create_sheet("Стили")
t.append(["основной стиль", "в профиле", "позиций", "средний балл", "если 1 330: оплатить", "если 1 330: тест", "сумма ¥", "главные сегменты"])
from collections import defaultdict
PROF = {"скульптурное золото", "жемчуг", "чёрный акцент", "цветной камень в ореоле", "бант / сердце", "змея / омега", "праздничные стразы"}
by = defaultdict(list)
for r in rows: by[r["style"].split(" (")[0]].append(r)
for k, v in sorted(by.items(), key=lambda x: -len(x[1])):
    t.append([k, "да" if k in PROF else "нет", len(v), round(sum(r["total"] for r in v) / len(v), 1), sum(r["decision1330"] == "оплатить" for r in v), sum(r["decision1330"] == "тест 20–30 шт" for r in v),
              round(sum(float(r["cny"]) * float(r["qty"]) for r in v)), ", ".join(f"{a} {b}" for a, b in Counter(r["seg"] for r in v).most_common(3))])
for c in t[1]: c.font = Font(bold=True)
for col, w in zip("ABCDEFGH", [26, 9, 9, 9, 10, 10, 10, 60]): t.column_dimensions[col].width = w
g = wb.create_sheet("Справочник стилей")
for line in open("стили/СТИЛИ.md").read().split("\n"): g.append([line])
g.column_dimensions["A"].width = 200
m = wb.create_sheet("Метод")
for line in open("METHOD.txt").read().split("\n"): m.append([line])
m.column_dimensions["A"].width = 160
wb.save(OUT); print(OUT)
